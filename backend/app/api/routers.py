"""API 路由（控制器）：解析请求、调用服务、格式化响应，不写业务逻辑。"""
import csv
import io
import json
import zipfile
from datetime import date, datetime, timedelta, timezone
from typing import Annotated

from fastapi import APIRouter, Depends, File, Form, Query, UploadFile
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse, Response
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import (
    get_current_permissions,
    get_current_user,
    require_permissions,
    require_roles,
)
from app.core.db import get_db
from app.core.errors import ForbiddenError, NotFoundError, ValidationError
from app.realtime.events import get_event_bus
from app.models import User
from app.schemas import (
    AnnouncementCreate,
    AnnouncementOut,
    AnnouncementReadersOut,
    AnnouncementReadOut,
    AnnouncementUpdate,
    AttendanceAlertNotifyIn,
    AttendanceAlertNotifyOut,
    AttendanceAlertsOut,
    AttendanceClockOut,
    AttendanceImportResult,
    AttendanceMeOut,
    AttendanceRecordOut,
    AttendanceRecordUpdate,
    AttendanceSettingIn,
    AttendanceSettingOut,
    AttendanceStatsOut,
    BatchPayRequest,
    BatchPayResponse,
    BudgetCreate,
    BudgetOut,
    BudgetPage,
    BudgetQuotaCheck,
    BudgetUpdate,
    DashboardOut,
    DelegationCreate,
    DelegationOut,
    DelegationUpdate,
    DepartmentCreate,
    DepartmentOut,
    DepartmentUpdate,
    DocumentOut,
    DocumentUpdate,
    DocumentVersionOut,
    DocumentsPage,
    ExpenseCreate,
    ExpenseOut,
    ExpensePage,
    ExpensePaidIn,
    ExpenseStatsOut,
    ExpenseUpdate,
    FieldWorkCreate,
    FieldWorkOut,
    FieldWorkPage,
    FieldWorkReview,
    FormDefinitionCreate,
    FormDefinitionOut,
    InvoiceVerifyRequest,
    InvoiceVerifyResponse,
    LeaveCreate,
    LeaveOut,
    LeaveReview,
    LoginRequest,
    MeOut,
    MenuCreate,
    MenuOut,
    MenuUpdate,
    Message,
    ModelEntityCreate,
    ModelEntityOut,
    ModelEntityUpdate,
    ModelRecordCreate,
    ModelRecordOut,
    ModelRecordsPage,
    ModelRecordUpdate,
    ModelStatsOut,
    OfficialDocCirculate,
    OfficialDocCreate,
    OfficialDocOut,
    OfficialDocPage,
    OfficialDocUpdate,
    OvertimeCreate,
    OvertimeOut,
    OvertimeReview,
    PermissionAssign,
    PermissionCreate,
    PermissionOut,
    PortalOverviewOut,
    RefreshRequest,
    RoleCreate,
    RoleOut,
    RoleUpdate,
    SignatureCreate,
    SignatureOut,
    TokenResponse,
    UserCreate,
    UserOut,
    UserUpdate,
    WorkflowAddSign,
    WorkflowDefinitionCreate,
    WorkflowDefinitionOut,
    WorkflowInstanceCreate,
    WorkflowInstanceOut,
    WorkflowReview,
    WorkflowTaskOut,
    WorkflowTransfer,
    WorkflowUrge,
    WorkflowStatsOut,
    WorkflowVersionOut,
    WorkflowVersionSave,
    OpsOverviewOut,
    SchedulerRunOut,
    AuditLogsPage,
    AuditStatsOut,
    AuditPurgeRequest,
    AuditPurgeOut,
    AuditLogOut,
    AuditLogsByCorrelationOut,
    AuditDurationStatsOut,
    NotificationOut,
    NotificationsPage,
)
from app.services import (
    AnnouncementService,
    AttendanceService,
    AuditService,
    AuthService,
    BatchPaymentService,
    BudgetService,
    DashboardService,
    DelegationService,
    DepartmentService,
    DocumentService,
    ExpenseService,
    InvoiceService,
    LeaveService,
    MenuService,
    ModelingService,
    NotificationService,
    OfficialDocService,
    OpsService,
    PermissionService,
    PortalService,
    ReportService,
    RoleService,
    SignatureService,
    UserService,
    WorkflowService,
    _to_utc_naive,
)

Db = Annotated[AsyncSession, Depends(get_db)]
CurrentUser = Annotated[User, Depends(get_current_user)]
AdminUser = Annotated[User, Depends(require_roles("admin"))]
ManagerUser = Annotated[User, Depends(require_roles("admin", "manager"))]
RoleViewDep = Annotated[None, Depends(require_permissions("org:role:view"))]
RoleManageDep = Annotated[None, Depends(require_permissions("org:role:manage"))]
# 当前用户权限码（用于「可选可见」场景，如报表中心组织级视图）
CurrentPerms = Annotated[list[str], Depends(get_current_permissions)]


def _wf(session: AsyncSession) -> WorkflowService:
    """构造工作流服务并注入进程级事件总线，使审批动作可实时推送。"""
    return WorkflowService(session, event_bus=get_event_bus())


def _ann(session: AsyncSession) -> AnnouncementService:
    """公告服务：注入事件总线，使发布公告可实时广播。"""
    return AnnouncementService(session, event_bus=get_event_bus())


def _leave(session: AsyncSession) -> LeaveService:
    """请假服务：注入事件总线，使请假待办可实时通知审批人。"""
    return LeaveService(session, event_bus=get_event_bus())


def _delegation(session: AsyncSession) -> DelegationService:
    """委托服务：注入事件总线，使委托创建/更新/删除可实时通知受托人。"""
    return DelegationService(session, event_bus=get_event_bus())


def _att(session: AsyncSession) -> AttendanceService:
    """考勤服务：注入事件总线，使外勤审批与异常预警可实时推送。"""
    return AttendanceService(session, event_bus=get_event_bus())


def _with_role_name(user: User) -> User:
    """为响应注入角色名，避免前端用 role_id 硬编码判断权限。"""
    user.role_name = user.role.name
    return user


# 容器路由：只负责统一挂载，**不要**再加 prefix——
# 各子路由已自带 /api/... 前缀（auth_router="/api/auth" 等），若此处再加 "/api"，
# include_router 会拼成 /api/api/...，导致前端请求的 /api/** 全部 404。
api = APIRouter()

# ---------- 认证 ----------
auth_router = APIRouter(prefix="/api/auth", tags=["auth"])


@auth_router.post("/login", response_model=TokenResponse)
async def login(body: LoginRequest, session: Db):
    access, refresh, _ = await AuthService(session).login(body.username, body.password)
    return TokenResponse(access_token=access, refresh_token=refresh)


@auth_router.post("/refresh", response_model=TokenResponse)
async def refresh(body: RefreshRequest, session: Db):
    access, refresh = await AuthService(session).refresh(body.refresh_token)
    return TokenResponse(access_token=access, refresh_token=refresh)


@auth_router.post("/logout", response_model=Message)
async def logout(body: RefreshRequest, session: Db):
    await AuthService(session).logout(body.refresh_token)
    return Message(message="已退出登录")


@auth_router.get("/me", response_model=MeOut)
async def me(user: CurrentUser, session: Db):
    """返回用户、角色、权限码、数据范围，前端据此动态渲染菜单与按钮权限。"""
    svc = PermissionService(session)
    codes = await svc.role_permission_codes(user.role_id)
    if user.role.name == "admin":
        codes = ["*"]
    return MeOut(
        id=user.id,
        username=user.username,
        full_name=user.full_name,
        email=user.email,
        is_active=user.is_active,
        role_id=user.role_id,
        role_name=user.role.name,
        department_id=user.department_id,
        created_at=user.created_at,
        permissions=codes,
        data_scope=user.role.data_scope,
    )


@auth_router.get("/menus", response_model=list[MenuOut])
async def my_menus(user: CurrentUser, session: Db):
    """当前用户可见菜单（按权限过滤），前端据此动态渲染导航。"""
    svc = PermissionService(session)
    codes = await svc.role_permission_codes(user.role_id)
    if user.role.name == "admin":
        codes = ["*"]
    menus = await MenuService(session).list_visible(codes)
    return [MenuOut(**m) for m in menus]


# ---------- 用户 / 角色 ----------
user_router = APIRouter(prefix="/api/users", tags=["users"])


@user_router.get("", response_model=list[UserOut])
async def list_users(
    session: Db, _: CurrentUser, skip: int = 0, limit: int = 200, q: str | None = None
):
    users = await UserService(session).list(skip, limit, q)
    return [_with_role_name(u) for u in users]


@user_router.post("", response_model=UserOut)
async def create_user(body: UserCreate, session: Db, _: AdminUser):
    return _with_role_name(await UserService(session).create(body))


@user_router.get("/roles", response_model=list[RoleOut])
async def list_roles(session: Db, _: RoleViewDep):
    """角色列表；需 org:role:view 权限（管理员/被授权者可见）。"""
    return await RoleService(session).list()


@user_router.post("/roles", response_model=RoleOut, status_code=201)
async def create_role(body: RoleCreate, session: Db, _: RoleManageDep):
    """新建角色；name 为唯一英文标识，服务层校验重名。"""
    return await RoleService(session).create(body)


@user_router.patch("/roles/{role_id}", response_model=RoleOut)
async def update_role(role_id: int, body: RoleUpdate, session: Db, _: RoleManageDep):
    """更新角色描述与数据范围（角色名作为稳定标识不可改）。"""
    return await RoleService(session).update(role_id, body)


@user_router.delete("/roles/{role_id}", response_model=Message)
async def delete_role(role_id: int, session: Db, _: RoleManageDep):
    """删除角色；admin 不可删、名下仍有用户不可删（护栏在服务层）。"""
    await RoleService(session).delete(role_id)
    return Message(message="角色已删除")


@user_router.get("/{user_id}", response_model=UserOut)
async def get_user(user_id: int, session: Db, _: CurrentUser):
    return _with_role_name(await UserService(session).get(user_id))


@user_router.patch("/{user_id}", response_model=UserOut)
async def update_user(user_id: int, body: UserUpdate, session: Db, user: CurrentUser):
    # 越权判断下沉到服务层（服务层按 actor 角色剥离提权字段）
    return _with_role_name(await UserService(session).update(user_id, body, user))


# ---------- 部门 ----------
dept_router = APIRouter(prefix="/api/departments", tags=["departments"])


@dept_router.get("", response_model=list[DepartmentOut])
async def list_depts(session: Db, _: CurrentUser):
    return await DepartmentService(session).list()


@dept_router.post("", response_model=DepartmentOut)
async def create_dept(body: DepartmentCreate, session: Db, _: AdminUser):
    return await DepartmentService(session).create(body)


@dept_router.patch("/{dept_id}", response_model=DepartmentOut)
async def update_dept(dept_id: int, body: DepartmentUpdate, session: Db, _: AdminUser):
    return await DepartmentService(session).update(dept_id, body)


# ---------- 考勤 / 请假 ----------
leave_router = APIRouter(prefix="/api/leave", tags=["leave"])


@leave_router.get("", response_model=list[LeaveOut])
async def list_leave(session: Db, user: CurrentUser, skip: int = 0, limit: int = 50):
    # 数据范围：all=全部 / dept=本部门 / self=本人
    scope = user.role.data_scope
    uid = user.id if scope == "self" else None
    dept = user.department_id if scope == "dept" else None
    return await _leave(session).list(uid, dept)


@leave_router.post("", response_model=LeaveOut)
async def create_leave(body: LeaveCreate, session: Db, user: CurrentUser):
    return await _leave(session).create(user.id, body)


@leave_router.post("/{leave_id}/review", response_model=LeaveOut)
async def review_leave(leave_id: int, body: LeaveReview, session: Db, _: ManagerUser):
    return await _leave(session).review(leave_id, _, body)


# ---------- 公告 / 文档 ----------
ann_router = APIRouter(prefix="/api/announcements", tags=["announcements"])


@ann_router.get("", response_model=list[AnnouncementOut])
async def list_anns(session: Db, user: CurrentUser, published_only: bool = True):
    # 数据范围=all（管理员）可见未发布公告，其余仅见已发布
    only = published_only if user.role.data_scope != "all" else False
    return await _ann(session).list(only)


@ann_router.post("", response_model=AnnouncementOut)
async def create_ann(body: AnnouncementCreate, session: Db, user: ManagerUser):
    return await _ann(session).create(user.id, body)


@ann_router.patch("/{ann_id}", response_model=AnnouncementOut)
async def update_ann(ann_id: int, body: AnnouncementUpdate, session: Db, _: ManagerUser):
    return await _ann(session).update(ann_id, body)


# ---------- 审批工作流 ----------
wf_router = APIRouter(prefix="/api/workflows", tags=["workflows"])


@wf_router.get("/definitions", response_model=list[WorkflowDefinitionOut])
async def list_defs(session: Db, _: CurrentUser, status: str | None = None):
    """流程列表；传 status=published 可只取可发起的流程。"""
    return await _wf(session).list_definitions(status)


@wf_router.post("/definitions", response_model=WorkflowDefinitionOut)
async def create_def(body: WorkflowDefinitionCreate, session: Db, _: AdminUser):
    return await _wf(session).create_definition(body)


@wf_router.patch("/definitions/{def_id}/status", response_model=WorkflowDefinitionOut)
async def set_def_status(
    def_id: int, session: Db, _: AdminUser, status: str = Query(..., pattern="^(draft|published|locked)$")
):
    """发布 / 锁定。锁定后不可发起，但在途实例继续流转。"""
    return await _wf(session).set_definition_status(def_id, status)


# ---- 表单设计 ----
@wf_router.get("/forms", response_model=list[FormDefinitionOut])
async def list_forms(session: Db, _: CurrentUser):
    return await _wf(session).list_forms()


@wf_router.post("/forms", response_model=FormDefinitionOut)
async def create_form(body: FormDefinitionCreate, session: Db, _: AdminUser):
    return await _wf(session).create_form(body)


# ---- 版本与流程图 ----
@wf_router.post("/definitions/{def_id}/versions", response_model=WorkflowVersionOut)
async def save_version(def_id: int, body: WorkflowVersionSave, session: Db, _: AdminUser):
    return await _wf(session).save_version(def_id, body)


@wf_router.get("/definitions/{def_id}/versions", response_model=list[WorkflowVersionOut])
async def list_versions(def_id: int, session: Db, _: CurrentUser):
    return await _wf(session).list_versions(def_id)


@wf_router.post("/versions/{version_id}/publish", response_model=WorkflowVersionOut)
async def publish_version(version_id: int, session: Db, _: AdminUser):
    return await _wf(session).publish_version(version_id)


# ---- 实例 ----
@wf_router.post("/instances", response_model=WorkflowInstanceOut)
async def start_instance(body: WorkflowInstanceCreate, session: Db, user: CurrentUser):
    return await _wf(session).start_instance(user.id, body)


@wf_router.get("/instances", response_model=list[WorkflowInstanceOut])
async def list_instances(session: Db, user: CurrentUser, status: str | None = None):
    uid = None if user.role.name == "admin" else user.id
    return await _wf(session).list_instances(uid, status)


@wf_router.get("/instances/{instance_id}", response_model=WorkflowInstanceOut)
async def get_instance(instance_id: int, session: Db, _: CurrentUser):
    return await _wf(session).get_instance(instance_id)


@wf_router.get("/instances/{instance_id}/logs", response_model=list[dict])
async def instance_logs(instance_id: int, session: Db, _: CurrentUser):
    logs = await _wf(session).list_logs(instance_id)
    return [
        {
            "id": l.id,
            "node_id": l.node_id,
            "node_name": l.node_name,
            "action": l.action,
            "operator_id": l.operator_id,
            "comment": l.comment,
            "created_at": l.created_at,
        }
        for l in logs
    ]


# ---- 流程中心四个视图 ----
@wf_router.get("/pending", response_model=list[WorkflowTaskOut])
async def my_pending(session: Db, user: CurrentUser):
    return await _wf(session).pending_for(user)


@wf_router.get("/done", response_model=list[WorkflowTaskOut])
async def my_done(session: Db, user: CurrentUser):
    return await _wf(session).done_for(user)


@wf_router.get("/notified", response_model=list[WorkflowTaskOut])
async def my_notified(session: Db, user: CurrentUser):
    return await _wf(session).notified_for(user)


@wf_router.get("/mine", response_model=list[WorkflowInstanceOut])
async def my_instances(session: Db, user: CurrentUser):
    return await _wf(session).my_instances(user)


# ---- 流程监控看板 ----
WfMonitorDep = Annotated[None, Depends(require_permissions("wf:monitor"))]


@wf_router.get("/stats", response_model=WorkflowStatsOut)
async def workflow_stats(session: Db, _: WfMonitorDep):
    """流程监控看板聚合数据；需 wf:monitor 权限（管理员/经理可见）。"""
    return await _wf(session).get_stats()


# ---- 审批动作 ----
@wf_router.post("/tasks/{task_id}/decide", response_model=WorkflowInstanceOut)
async def decide_task(task_id: int, body: WorkflowReview, session: Db, user: CurrentUser):
    return await _wf(session).decide_task(task_id, user, body)


@wf_router.post("/tasks/{task_id}/transfer", response_model=WorkflowInstanceOut)
async def transfer_task(task_id: int, body: WorkflowTransfer, session: Db, user: CurrentUser):
    return await _wf(session).transfer_task(task_id, user, body)


@wf_router.post("/instances/{instance_id}/nodes/{node_id}/addsign", response_model=WorkflowInstanceOut)
async def addsign(
    instance_id: int, node_id: str, body: WorkflowAddSign, session: Db, user: CurrentUser
):
    return await _wf(session).addsign(instance_id, node_id, user, body)


@wf_router.post("/instances/{instance_id}/withdraw", response_model=WorkflowInstanceOut)
async def withdraw(instance_id: int, session: Db, user: CurrentUser):
    return await _wf(session).withdraw(instance_id, user)


@wf_router.post("/instances/{instance_id}/urge", response_model=dict)
async def urge(instance_id: int, body: WorkflowUrge, session: Db, user: CurrentUser):
    count = await _wf(session).urge(instance_id, user, body)
    return {"urged": count}


# ---------- 电子签单 ----------
sig_router = APIRouter(prefix="/api/signatures", tags=["signatures"])


@sig_router.get("/{ref_type}/{ref_id}", response_model=list[SignatureOut])
async def list_signatures(ref_type: str, ref_id: int, session: Db, _: CurrentUser):
    return await SignatureService(session).list_by_ref(ref_type, ref_id)


@sig_router.post("", response_model=SignatureOut)
async def create_signature(body: SignatureCreate, session: Db, user: CurrentUser):
    return await SignatureService(session).create(user.id, body)


# ---------- 组织权限中心：权限 / 菜单 ----------
perm_router = APIRouter(prefix="/api/permissions", tags=["permissions"])

PermManageDep = Annotated[None, Depends(require_permissions("org:role:manage"))]
PermViewDep = Annotated[None, Depends(require_permissions("org:role:view"))]


@perm_router.get("", response_model=list[PermissionOut])
async def list_permissions(
    session: Db, _: CurrentUser, category: str | None = None, ptype: str | None = None
):
    return await PermissionService(session).list(category, ptype)


@perm_router.post("", response_model=PermissionOut)
async def create_permission(body: PermissionCreate, session: Db, _: PermManageDep):
    return await PermissionService(session).create(body)


@perm_router.get("/roles/{role_id}", response_model=list[PermissionOut])
async def role_permissions(role_id: int, session: Db, _: PermViewDep):
    return await PermissionService(session).list_role_permissions(role_id)


@perm_router.post("/roles/{role_id}", response_model=Message)
async def assign_permission(role_id: int, body: PermissionAssign, session: Db, _: PermManageDep):
    await PermissionService(session).assign(role_id, body.permission_id)
    return Message(message="已授权")


@perm_router.delete("/roles/{role_id}/{perm_id}", response_model=Message)
async def revoke_permission(role_id: int, perm_id: int, session: Db, _: PermManageDep):
    await PermissionService(session).revoke(role_id, perm_id)
    return Message(message="已撤销")


menu_router = APIRouter(prefix="/api/menus", tags=["menus"])

MenuManageDep = Annotated[None, Depends(require_permissions("org:menu:manage"))]
MenuViewDep = Annotated[None, Depends(require_permissions("org:menu:view"))]


@menu_router.get("", response_model=list[MenuOut])
async def list_menus(session: Db, _: MenuViewDep):
    return [MenuOut(**m) for m in await MenuService(session).list_all()]


@menu_router.post("", response_model=MenuOut)
async def create_menu(body: MenuCreate, session: Db, _: MenuManageDep):
    return await MenuService(session).create(body)


@menu_router.patch("/{menu_id}", response_model=MenuOut)
async def update_menu(menu_id: int, body: MenuUpdate, session: Db, _: MenuManageDep):
    return await MenuService(session).update(menu_id, body)


# ---------- 内容门户：角色化工作台 ----------
portal_router = APIRouter(prefix="/api/portal", tags=["portal"])


@portal_router.get("/overview", response_model=PortalOverviewOut)
async def portal_overview(session: Db, user: CurrentUser):
    """工作台门户聚合：按角色返回统计与待办/在途/请假/公告挂件，供首页渲染。"""
    return await PortalService(session).overview(user)


# ---------- 运维中心（P3 建模运维） ----------
ops_router = APIRouter(prefix="/api/ops", tags=["ops"])


@ops_router.get("/overview", response_model=OpsOverviewOut)
async def ops_overview(session: Db, user: AdminUser):
    """运维中心：管理员视角的系统级运行概览（实体计数、流程分布、角色矩阵、最近实例）。"""
    return await OpsService(session).overview(user)


@ops_router.post("/scheduler/run-once", response_model=SchedulerRunOut)
async def run_scheduler_once(session: Db, user: AdminUser):
    """运维手动触发一轮审批超时升级（默认由后台调度器周期执行）。需管理员权限。"""
    return await _wf(session).timeout_escalate_all()


# ---------- 报表中心：统一数据看板 ----------
dashboard_router = APIRouter(prefix="/api/dashboard", tags=["dashboard"])


@dashboard_router.get("", response_model=DashboardOut)
async def dashboard_overview(
    session: Db,
    user: CurrentUser,
    perms: CurrentPerms,
    days: int = Query(30, ge=7, le=365, description="组织级趋势窗口天数（7/30/90）"),
):
    """报表中心聚合：个人视图恒返回，组织级 KPI/图表需 `report:view`（管理员通配放行）。

    `days` 仅影响组织级趋势分桶窗口；个人视图与 KPI 等全量聚合不受其影响。
    """
    has_org = "*" in perms or "report:view" in perms
    return await DashboardService(session).overview(user, has_org_scope=has_org, days=days)


@dashboard_router.get("/export")
async def dashboard_export(
    session: Db,
    _: Annotated[None, Depends(require_permissions("report:view"))],
    days: int = Query(30, ge=7, le=365, description="导出数据的时间窗口天数"),
    format: str = Query("csv", pattern="^(csv)$", description="导出格式，当前仅支持 csv（多表打包为 zip）"),
):
    """组织级报表导出：仅 `report:view` 可调用，返回多张 CSV 打包的 zip。

    导出内容与组织级看板一致（KPI / 流程状态 / 趋势 / 部门人数 / 最近动态），
    因含全公司聚合数据，故与看板同源受 report:view 约束。
    """
    tables = await DashboardService(session).export_tables(days)
    payload = _csv_zip(tables)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d")
    filename = f"zk-oa-dashboard-{days}d-{stamp}.zip"
    return Response(
        content=payload,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@dashboard_router.get("/cube")
async def report_cube(
    session: Db,
    _: Annotated[None, Depends(require_permissions("report:view"))],
    metrics: str = Query("attendance,expense,workflow,official_doc", description="逗号分隔的指标列表"),
    dimension: str = Query("month", description="聚合维度: month/user/department"),
    start: str | None = None,
    end: str | None = None,
):
    """多维数据立方体：按维度聚合指定指标（含全公司聚合，需 report:view）。"""
    from datetime import date as _date
    start_d = _date.fromisoformat(start) if start else None
    end_d = _date.fromisoformat(end) if end else None
    metric_list = [m.strip() for m in metrics.split(",") if m.strip()]
    return await ReportService(session).cube(metric_list, dimension, start_d, end_d)


@dashboard_router.get("/trends")
async def report_trends(
    session: Db,
    _: Annotated[None, Depends(require_permissions("report:view"))],
    days: int = Query(30, ge=7, le=365),
    metrics: str | None = None,
):
    """趋势分析：近 N 天各指标的日趋势 + 7天移动平均（含全公司聚合，需 report:view）。"""
    metric_list = [m.strip() for m in metrics.split(",") if m.strip()] if metrics else None
    return await ReportService(session).trends(days, metric_list)


def _csv_zip(tables: dict[str, list[list[str]]]) -> bytes:
    """把多张二维字符串表编码为带 BOM 的 ZIP（Excel 可直接打开，中文不乱码）。

    纯标准库实现，无第三方依赖；服务层只产出 `tables`，字节编码在此完成。
    """
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, rows in tables.items():
            sio = io.StringIO()
            csv.writer(sio).writerows(rows)
            # 前置 UTF-8 BOM，保证 Windows Excel 正确识别中文
            zf.writestr(name, "\ufeff" + sio.getvalue())
    return buf.getvalue()


# ---------- 日志中心（P3 建模运维）：审计日志 ----------
audit_router = APIRouter(prefix="/api/audit", tags=["audit"])

AuditViewDep = Annotated[None, Depends(require_permissions("audit:view"))]
AuditManageDep = Annotated[None, Depends(require_permissions("audit:manage"))]


def _parse_dt(s: str | None) -> datetime | None:
    """ISO 时间窗解析；格式非法时返回 None（不抛出，避免泄露细节）。"""
    if not s:
        return None
    try:
        return datetime.fromisoformat(s)
    except ValueError:
        return None


@audit_router.get("/logs", response_model=AuditLogsPage)
async def audit_logs(
    session: Db, _: AuditViewDep,
    actor_id: int | None = None, module: str | None = None,
    action: str | None = None, status: str | None = None,
    q: str | None = None, start: str | None = None, end: str | None = None,
    offset: int = 0, limit: int = 50,
):
    """审计日志列表（按权限可见）；支持模块/动作/状态/关键字/时间窗过滤。"""
    return await AuditService(session).list_logs(
        actor_id=actor_id, module=module, action=action, status=status,
        q=q, start=_parse_dt(start), end=_parse_dt(end), offset=offset, limit=limit,
    )


@audit_router.get("/stats", response_model=AuditStatsOut)
async def audit_stats(session: Db, _: AuditViewDep):
    """日志中心概览：总量、成功/失败、按模块/状态分布、近 14 天趋势。"""
    return await AuditService(session).stats()


@audit_router.post("/purge", response_model=AuditPurgeOut)
async def audit_purge(body: AuditPurgeRequest, session: Db, _: AuditManageDep):
    """清理 N 天前的审计日志（仅 audit:manage）。"""
    deleted = await AuditService(session).purge(body.days)
    before = datetime.now(timezone.utc) - timedelta(days=body.days)
    return AuditPurgeOut(deleted=deleted, before=before)


@audit_router.get("/logs/{log_id}", response_model=AuditLogOut)
async def audit_log_detail(log_id: int, session: Db, _: AuditViewDep):
    """获取单条审计日志详情（含 before/after value, correlation_id, duration_ms）。"""
    detail = await AuditService(session).get_log_detail(log_id)
    if not detail:
        raise NotFoundError("审计日志不存在")
    return AuditLogOut(**detail)


@audit_router.get("/logs/correlation/{correlation_id}", response_model=AuditLogsByCorrelationOut)
async def audit_logs_by_correlation(
    correlation_id: str,
    session: Db,
    _: AuditViewDep,
    offset: int = 0,
    limit: int = 50,
):
    """按关联 ID 查询关联的审计日志"""
    return await AuditService(session).list_logs_by_correlation(correlation_id, offset, limit)


@audit_router.get("/stats/duration", response_model=AuditDurationStatsOut)
async def audit_stats_duration(
    session: Db, _: AuditViewDep, module: str | None = None
):
    """统计请求耗时分布（可选按模块筛选）"""
    return await AuditService(session).stats_duration(module)


@audit_router.get("/export")
async def audit_export(
    session: Db, _: AuditViewDep,
    module: str | None = None,
    action: str | None = None,
    status: str | None = None,
    start: str | None = None,
    end: str | None = None,
    limit: int = Query(10000, ge=1, le=50000),
):
    """导出审计日志为 CSV（供下载分析）。"""
    import csv
    import io
    from fastapi.responses import StreamingResponse

    start_dt = datetime.fromisoformat(start) if start else None
    end_dt = datetime.fromisoformat(end) if end else None
    data = await AuditService(session).export_logs(
        module=module, action=action, status=status,
        start=start_dt, end=end_dt, limit=limit,
    )
    if not data:
        return {"message": "无数据可导出", "count": 0}
    # 生成 CSV
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=data[0].keys())
    writer.writeheader()
    writer.writerows(data)
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=audit_logs_{datetime.now():%Y%m%d_%H%M%S}.csv"},
    )


# ---------- 建模引擎（P3）：自定义实体 + 动态数据 ----------
model_router = APIRouter(prefix="/api/models", tags=["models"])

ModelViewDep = Annotated[None, Depends(require_permissions("model:view"))]
ModelManageDep = Annotated[None, Depends(require_permissions("model:manage"))]
ModelDataDep = Annotated[None, Depends(require_permissions("model:data:edit"))]


def _parse_filters(raw: str | None) -> dict | None:
    """过滤参数为 JSON 对象字符串，格式错误统一转 422，不让异常细节外泄。"""
    if not raw:
        return None
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        raise ValidationError("过滤参数格式错误，需为 JSON 对象")
    if not isinstance(parsed, dict):
        raise ValidationError("过滤参数格式错误，需为 JSON 对象")
    return parsed


@model_router.get("/entities", response_model=list[ModelEntityOut])
async def list_model_entities(session: Db, _: ModelViewDep, status: str | None = None):
    """实体列表（含记录数），按 model:view 权限可见。"""
    return await ModelingService(session).list_entities(status)


@model_router.post("/entities", response_model=ModelEntityOut, status_code=201)
async def create_model_entity(body: ModelEntityCreate, session: Db, user: CurrentUser,
                             _: ModelManageDep):
    return await ModelingService(session).create_entity(body, user)


@model_router.get("/entities/{entity_id}", response_model=ModelEntityOut)
async def get_model_entity(entity_id: int, session: Db, _: ModelViewDep):
    return await ModelingService(session).get_entity(entity_id)


@model_router.patch("/entities/{entity_id}", response_model=ModelEntityOut)
async def update_model_entity(
    entity_id: int, body: ModelEntityUpdate, session: Db, _: ModelManageDep
):
    return await ModelingService(session).update_entity(entity_id, body)


@model_router.delete("/entities/{entity_id}", response_model=Message)
async def delete_model_entity(entity_id: int, session: Db, _: ModelManageDep):
    await ModelingService(session).delete_entity(entity_id)
    return Message(message="实体已删除")


@model_router.get("/entities/{entity_id}/records", response_model=ModelRecordsPage)
async def list_model_records(
    entity_id: int, session: Db, _: ModelViewDep,
    filters: str | None = None, q: str | None = None,
    offset: int = 0, limit: int = 50,
):
    """实体数据列表：filters 为 JSON 字符串（{"status":"启用"}），q 为全字段模糊搜索。"""
    return await ModelingService(session).list_records(
        entity_id, filters=_parse_filters(filters), q=q, offset=offset, limit=limit
    )


@model_router.post("/entities/{entity_id}/records", response_model=ModelRecordOut,
                   status_code=201)
async def create_model_record(
    entity_id: int, body: ModelRecordCreate, session: Db, user: CurrentUser, _: ModelDataDep
):
    """新增数据行：仅已发布实体可录入，字段级校验在服务层。"""
    return await ModelingService(session).create_record(entity_id, body, user)


@model_router.patch("/records/{record_id}", response_model=ModelRecordOut)
async def update_model_record(
    record_id: int, body: ModelRecordUpdate, session: Db, user: CurrentUser, _: ModelDataDep
):
    return await ModelingService(session).update_record(record_id, body, user)


@model_router.delete("/records/{record_id}", response_model=Message)
async def delete_model_record(record_id: int, session: Db, _: ModelDataDep):
    await ModelingService(session).delete_record(record_id)
    return Message(message="数据已删除")


@model_router.get("/stats", response_model=ModelStatsOut)
async def model_stats(session: Db, _: ModelViewDep):
    """建模引擎概览：实体数 / 数据量 / 今日新增 / Top 实体。"""
    return await ModelingService(session).stats()


# ---------- 内容引擎：文档中心 + 文件上传 ----------
doc_router = APIRouter(prefix="/api/documents", tags=["documents"])

DocViewDep = Annotated[None, Depends(require_permissions("doc:view"))]
DocUploadDep = Annotated[None, Depends(require_permissions("doc:upload"))]
DocManageDep = Annotated[None, Depends(require_permissions("doc:manage"))]


@doc_router.get("", response_model=DocumentsPage)
async def list_docs(
    session: Db, _: DocViewDep,
    category: str | None = None, q: str | None = None,
    offset: int = 0, limit: int = 50,
):
    """文档列表：按分类过滤 + 标题搜索，检索直接下推数据库。"""
    return await DocumentService(session).list_documents(
        category=category, q=q, offset=offset, limit=limit
    )


@doc_router.post("", response_model=DocumentOut, status_code=201)
async def upload_doc(
    session: Db, user: CurrentUser, _: DocUploadDep,
    title: str = Form(...), category: str = Form("general"),
    description: str = Form(""), file: UploadFile = File(...),
):
    """上传并创建文档：文件以 {id}.{ext} 落盘，扩展名受白名单约束，大小受上限约束。"""
    return await DocumentService(session).upload_and_create(
        user, title, category, description, file
    )


@doc_router.patch("/{doc_id}", response_model=DocumentOut)
async def update_doc(doc_id: int, body: DocumentUpdate, session: Db, _: DocUploadDep):
    """更新文档元数据（标题/分类/描述）；文件本身需重新上传。"""
    return await DocumentService(session).update(doc_id, body)


@doc_router.get("/{doc_id}/file")
async def download_doc(doc_id: int, session: Db, _: DocViewDep):
    """鉴权后返回文件：本地后端直接流式返回；S3 后端下发签名 URL 并 302 跳转。"""
    local_path, signed_url, download_name = await DocumentService(session).get_file(doc_id)
    if signed_url:
        return RedirectResponse(url=signed_url, status_code=302)
    return FileResponse(local_path, filename=download_name)


@doc_router.delete("/{doc_id}", response_model=Message)
async def delete_doc(doc_id: int, session: Db, _: DocManageDep):
    """删除文档：同时清理磁盘文件（best-effort）。"""
    await DocumentService(session).delete(doc_id)
    return Message(message="文档已删除")


# ---------- 消息中心：通知收件箱 ----------
msg_router = APIRouter(prefix="/api/notifications", tags=["notifications"])

MsgViewDep = Annotated[None, Depends(require_permissions("msg:view"))]


@msg_router.get("", response_model=NotificationsPage)
async def list_notifications(
    session: Db, user: CurrentUser, _: MsgViewDep,
    unread_only: bool = Query(False), limit: int = 50, offset: int = 0,
):
    """当前用户的通知列表（含未读总数）；unread_only=true 仅取未读。"""
    return await NotificationService(session).list_for_user(
        user.id, only_unread=unread_only, limit=limit, offset=offset
    )


@msg_router.post("/{nid}/read", response_model=Message)
async def mark_read(nid: int, session: Db, user: CurrentUser, _: MsgViewDep):
    """标记单条为已读（越权访问别人的通知会 404）。"""
    await NotificationService(session).mark_read(nid, user.id)
    return Message(message="已标记已读")


@msg_router.post("/read-all", response_model=Message)
async def mark_all_read(session: Db, user: CurrentUser, _: MsgViewDep):
    """全部标记为已读。"""
    await NotificationService(session).mark_all_read(user.id)
    return Message(message="已全部标记为已读")


# =========================================================================
# ---------- Excel 导入 / 导出（openpyxl，失败时给出可执行的提示） ----------
# =========================================================================
def _xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - 取决于运行环境
        raise ValidationError("导出 Excel 需要 openpyxl，请执行 pip install openpyxl") from exc
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title=title[:31])
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_response(filename: str, sheets: dict[str, list[list]]) -> Response:
    return Response(
        content=_xlsx_bytes(sheets),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# =========================================================================
# ---------- 考勤管理 ----------
# =========================================================================
attendance_router = APIRouter(prefix="/api/attendance", tags=["attendance"])
AttViewDep = Annotated[None, Depends(require_permissions("att:view"))]
AttManageDep = Annotated[None, Depends(require_permissions("att:manage"))]

# 工作流查看（发起/管理自己的委托）与公告读写依赖
WfViewDep = Annotated[None, Depends(require_permissions("wf:view"))]
AnnViewDep = Annotated[None, Depends(require_permissions("ann:view"))]
AnnManageDep = Annotated[None, Depends(require_permissions("ann:manage"))]


def _att_out(rec) -> AttendanceRecordOut:
    user = getattr(rec, "user", None)
    return AttendanceRecordOut(
        id=rec.id,
        user_id=rec.user_id,
        user_name=(user.full_name or user.username) if user else "",
        day=rec.day,
        check_in=rec.check_in,
        check_out=rec.check_out,
        status=rec.status,
        late_minutes=rec.late_minutes,
        early_minutes=rec.early_minutes,
        work_minutes=rec.work_minutes,
        source=rec.source,
        note=rec.note,
    )


async def _user_display(session: AsyncSession, user_id: int | None) -> str:
    if not user_id:
        return ""
    res = await session.execute(
        select(User.username, User.full_name).where(User.id == user_id)
    )
    row = res.first()
    return (row[1] or row[0]) if row else f"#{user_id}"


@attendance_router.get("/settings", response_model=AttendanceSettingOut)
async def get_attendance_settings(session: Db, _: AttViewDep):
    """班次设置（未初始化时返回默认值）。"""
    return await AttendanceService(session).get_settings()


@attendance_router.put("/settings", response_model=AttendanceSettingOut)
async def save_attendance_settings(
    body: AttendanceSettingIn, session: Db, user: CurrentUser, _: AttManageDep
):
    """保存班次（工作日 / 上下班时间 / 迟到宽限）。"""
    return await AttendanceService(session).save_settings(
        user.id, body.model_dump(exclude_unset=True)
    )


@attendance_router.post("/clock", response_model=AttendanceClockOut)
async def clock_in_out(
    session: Db,
    user: CurrentUser,
    _: AttViewDep,
    # GPS参数（可选，用于外勤打卡）
    lat: float | None = Query(None),
    lng: float | None = Query(None),
    address: str | None = Query(None),
    is_field_work: bool = Query(False),
    field_work_address: str | None = Query(None),
    field_work_note: str | None = Query(None),
):
    """打卡：当日上午第一次打上班卡，第二次打下班卡；重复第三次报错。

    支持GPS定位打卡，可选参数：lat, lng, address, is_field_work, field_work_address, field_work_note
    """
    rec, message = await AttendanceService(session).clock(
        user, lat=lat, lng=lng, address=address,
        is_field_work=is_field_work,
        field_work_address=field_work_address,
        field_work_note=field_work_note,
    )
    return AttendanceClockOut(
        day=rec.day,
        check_in=rec.check_in,
        check_out=rec.check_out,
        status=rec.status,
        late_minutes=rec.late_minutes,
        early_minutes=rec.early_minutes,
        message=message,
        # GPS字段
        check_in_lat=rec.check_in_lat,
        check_in_lng=rec.check_in_lng,
        check_out_lat=rec.check_out_lat,
        check_out_lng=rec.check_out_lng,
        check_in_address=rec.check_in_address,
        check_out_address=rec.check_out_address,
        is_field_work=rec.is_field_work,
        field_work_address=rec.field_work_address,
        field_work_note=rec.field_work_note,
    )


@attendance_router.get("/me", response_model=AttendanceMeOut)
async def my_attendance(
    session: Db,
    user: CurrentUser,
    _: AttViewDep,
    month: str = Query(..., pattern=r"^\d{4}-\d{2}$"),
):
    """我的月度考勤：记录列表 + 汇总（应出勤/实出勤/迟到/早退/缺勤/加班）。"""
    data = await AttendanceService(session).my_month(user.id, month)
    return AttendanceMeOut(
        records=[_att_out(r) for r in data["records"]], summary=data["summary"]
    )


@attendance_router.get("/records")
async def list_attendance_records(
    session: Db,
    _: AttManageDep,
    start: date = Query(...),
    end: date = Query(...),
    user_id: int | None = None,
    status: str | None = None,
    offset: int = 0,
    limit: int = 50,
):
    """全部员工考勤记录（需 att:manage），按日期区间 + 状态筛选。"""
    items, total = await AttendanceService(session).list_records(
        start=start, end=end, user_id=user_id, status=status, offset=offset, limit=limit
    )
    return {
        "items": [_att_out(i) for i in items],
        "total": total,
        "offset": offset,
        "limit": limit,
    }


@attendance_router.patch("/records/{rec_id}", response_model=AttendanceRecordOut)
async def update_attendance_record(
    rec_id: int, body: AttendanceRecordUpdate, session: Db, _: AttManageDep
):
    """人工修正考勤（补打卡 / 标记请假、出差、远程 / 说明）。"""
    rec = await AttendanceService(session).update_record(
        rec_id, body.model_dump(exclude_unset=True)
    )
    return _att_out(rec)


@attendance_router.get("/stats", response_model=AttendanceStatsOut)
async def attendance_stats(session: Db, _: AttManageDep, month: str = Query(...)):
    """按人汇总当月出勤/迟到/早退/缺勤/加班，供管理端看板。"""
    return await AttendanceService(session).stats(month)


@attendance_router.get("/overtime", response_model=list[OvertimeOut])
async def list_overtime(
    session: Db,
    user: CurrentUser,
    _: AttViewDep,
    perms: Annotated[list[str], Depends(get_current_permissions)],
    status: str | None = None,
    all: bool = False,
    offset: int = 0,
    limit: int = 50,
):
    """加班申请列表：默认自己的；持 att:manage 且 all=true 看全部。"""
    can_all = all and ("*" in perms or "att:manage" in perms)
    svc = AttendanceService(session)
    items, _total = await svc.list_overtime(
        user_id=None if can_all else user.id, status=status, offset=offset, limit=limit
    )
    out: list[OvertimeOut] = []
    for item in items:
        row = OvertimeOut.model_validate(item)
        row.user_name = await _user_display(session, item.user_id)
        out.append(row)
    return out


@attendance_router.post("/overtime", response_model=OvertimeOut, status_code=201)
async def create_overtime(body: OvertimeCreate, session: Db, user: CurrentUser, _: AttViewDep):
    item = await AttendanceService(session).create_overtime(user.id, body.model_dump())
    row = OvertimeOut.model_validate(item)
    row.user_name = user.full_name or user.username
    return row


@attendance_router.post("/overtime/{ot_id}/review", response_model=OvertimeOut)
async def review_overtime(
    ot_id: int, body: OvertimeReview, session: Db, user: CurrentUser, _: AttManageDep
):
    item = await AttendanceService(session).review_overtime(ot_id, user, body.model_dump())
    row = OvertimeOut.model_validate(item)
    row.user_name = await _user_display(session, item.user_id)
    return row


@attendance_router.get("/export")
async def export_attendance(session: Db, _: AttManageDep, month: str = Query(...)):
    """导出当月考勤明细为 Excel（xlsx）。"""
    y, m = (int(x) for x in month.split("-"))
    svc = AttendanceService(session)
    items, _total = await svc.list_records(
        start=date(y, m, 1),
        end=date(y + (1 if m == 12 else 0), 1 if m == 12 else m + 1, 1),
        limit=5000,
    )
    rows = [["日期", "姓名", "状态", "上班", "下班", "迟到(分)", "早退(分)", "工时(分)", "来源", "备注"]]
    for r in items:
        rows.append([
            r.day.isoformat(),
            (r.user.full_name or r.user.username) if r.user else f"#{r.user_id}",
            r.status,
            r.check_in.strftime("%H:%M") if r.check_in else "",
            r.check_out.strftime("%H:%M") if r.check_out else "",
            r.late_minutes,
            r.early_minutes,
            r.work_minutes,
            r.source,
            r.note,
        ])
    return _xlsx_response(f"attendance-{month}.xlsx", {"考勤明细": rows})


@attendance_router.get("/import/template")
async def attendance_import_template(_: AttManageDep):
    """下载考勤批量导入模板（表头 + 示例行，字段说明见模板备注列）。"""
    rows = [
        ["用户ID", "用户名", "日期", "上班时间", "下班时间", "状态", "迟到分钟",
         "早退分钟", "工时分钟", "来源", "备注", "外勤", "外勤地址", "外勤备注"],
        [1, "admin", "2026-09-01", "08:55", "18:05", "normal", 0, 0, 550, "clock", "", "", "", ""],
        ["", "admin", "2026-09-02", "09:12", "18:00", "迟到", 12, 0, 528, "import",
         "示例：状态支持中文/英文", "", "", ""],
        ["", "admin", "2026-09-03", "", "", "remote", 0, 0, 0, "manual",
         "示例：远程办公", "否", "", ""],
        ["", "admin", "2026-09-04", "09:00", "18:00", "normal", 0, 0, 540, "clock",
         "示例：外勤打卡（有已通过外勤单时自动关联）", "是", "客户现场-中科园区", "项目驻场"],
    ]
    return _xlsx_response("attendance_import_template.xlsx", {"考勤导入": rows})


@attendance_router.post("/import", response_model=AttendanceImportResult)
async def import_attendance(
    session: Db,
    user: CurrentUser,
    _: AttManageDep,
    file: UploadFile = File(...),
    dry_run: bool = Query(False, description="true=只校验不写库（预检）"),
):
    """Excel 批量导入考勤记录。

    表头需包含：用户ID/用户名、日期、上班时间、下班时间、状态、迟到分钟、早退分钟、
    工时分钟、来源、备注，以及可选的外勤/外勤地址/外勤备注。
    `dry_run=true` 返回预检结果（created/updated 为「将要」发生的数量），不落库。
    """
    result = await _att(session).import_attendance(file, user.id, dry_run=dry_run)
    return AttendanceImportResult(**result)


# ---- 外勤申请（外勤关联） ----
@attendance_router.get("/field-work", response_model=FieldWorkPage)
async def list_field_work(
    session: Db,
    user: CurrentUser,
    perms: Annotated[list[str], Depends(get_current_permissions)],
    status: str | None = None,
    start: date | None = None,
    end: date | None = None,
    mine: bool = True,
    offset: int = 0,
    limit: int = Query(50, ge=1, le=200),
):
    """外勤申请列表：默认自己的；持 att:manage 且 mine=false 看全部。"""
    can_all = ("*" in perms or "att:manage" in perms) and not mine
    return await _att(session).list_field_work(
        user_id=None if can_all else user.id,
        status=status, start=start, end=end, offset=offset, limit=limit,
    )


@attendance_router.post("/field-work", response_model=FieldWorkOut, status_code=201)
async def create_field_work(body: FieldWorkCreate, session: Db, user: CurrentUser, _: AttViewDep):
    """登记外勤申请（区间 + 地点 + 事由），审批通过后打卡自动关联。"""
    item = await _att(session).create_field_work(user, body.model_dump())
    return _att(session)._fieldwork_out(item)


@attendance_router.post("/field-work/{oid}/review", response_model=FieldWorkOut)
async def review_field_work(
    oid: int, body: FieldWorkReview, session: Db, user: CurrentUser, _: AttManageDep
):
    """审批外勤申请（通过后考勤自动关联；驳回则打卡不关联）。"""
    item = await _att(session).review_field_work(oid, user, body.model_dump())
    return _att(session)._fieldwork_out(item)


# ---- 考勤异常预警 ----
@attendance_router.get("/alerts", response_model=AttendanceAlertsOut)
async def attendance_alerts(
    session: Db,
    _: AttManageDep,
    start: date = Query(..., description="扫描开始日期"),
    end: date = Query(..., description="扫描结束日期"),
    threshold: int = Query(3, ge=1, le=30, description="频繁/连续判定阈值"),
    limit: int = Query(200, ge=1, le=500),
):
    """考勤异常预警：缺勤/未打下班卡/迟到早退频繁/夜间打卡/外勤无单/连续缺勤。"""
    return await _att(session).alerts(start=start, end=end, threshold=threshold, limit=limit)


@attendance_router.post("/alerts/notify", response_model=AttendanceAlertNotifyOut)
async def notify_attendance_alerts(
    body: AttendanceAlertNotifyIn, session: Db, _: AttManageDep
):
    """把当前预警汇总站内信推送给全部具备 att:manage 的在岗用户。"""
    return await _att(session).notify_alerts(
        start=body.start, end=body.end, threshold=body.threshold, limit=body.limit
    )


# =========================================================================
# ---------- 审批委托（代审） ----------
# =========================================================================
@wf_router.get("/delegations", response_model=list[DelegationOut])
async def list_delegations(
    session: Db,
    user: CurrentUser,
    perms: Annotated[list[str], Depends(get_current_permissions)],
    all: bool = False,
):
    """委托列表：默认自己的；管理员或持 wf:monitor 且 all=true 可看全部。"""
    can_all = all and ("*" in perms or "wf:monitor" in perms)
    return await _delegation(session).list(mine=user, show_all=can_all)


@wf_router.post("/delegations", response_model=DelegationOut, status_code=201)
async def create_delegation(body: DelegationCreate, session: Db, user: CurrentUser, _: WfViewDep):
    """新建委托：把 [start_at, end_at] 内的审批权交给受托人（流程为空=全部流程）。"""
    return await _delegation(session).create(user, body.model_dump())


@wf_router.patch("/delegations/{did}", response_model=DelegationOut)
async def update_delegation(did: int, body: DelegationUpdate, session: Db, user: CurrentUser):
    return await _delegation(session).update(
        user, did, body.model_dump(exclude_unset=True)
    )


@wf_router.delete("/delegations/{did}", response_model=Message)
async def delete_delegation(did: int, session: Db, user: CurrentUser):
    await _delegation(session).delete(user, did)
    return Message(message="委托已删除")


@wf_router.post("/delegations/check-conflicts")
async def check_delegation_conflicts(
    body: DelegationCreate, session: Db, user: CurrentUser, _: WfViewDep
):
    """预检委托冲突：返回重叠/循环检测结果，不创建委托——前端可据此给出实时提示。"""
    # 统一成 UTC naive（与 create 一致）
    start_at = _to_utc_naive(body.start_at)
    end_at = _to_utc_naive(body.end_at)
    conflicts = await _delegation(session).check_conflicts(
        from_user_id=user.id,
        to_user_id=body.to_user_id,
        start_at=start_at,
        end_at=end_at,
        definition_id=body.definition_id,
    )
    return {"conflicts": conflicts, "blocking": len(conflicts) > 0}


# =========================================================================
# ---------- 报销单 ----------
# =========================================================================
expense_router = APIRouter(prefix="/api/expenses", tags=["expenses"])
ExpenseViewDep = Annotated[None, Depends(require_permissions("expense:view"))]
ExpenseManageDep = Annotated[None, Depends(require_permissions("expense:manage"))]


async def _expense_owned_check(
    session: AsyncSession, eid: int, user: User, perms: list[str]
):
    claim = await ExpenseService(session).get(eid)
    if claim.user_id != user.id and not ("*" in perms or "expense:manage" in perms):
        raise ForbiddenError("只能查看自己的报销单")
    return claim


@expense_router.get("", response_model=ExpensePage)
async def list_expenses(
    session: Db,
    user: CurrentUser,
    _: ExpenseViewDep,
    perms: Annotated[list[str], Depends(get_current_permissions)],
    status: str | None = None,
    month: str | None = None,
    q: str | None = None,
    mine: bool = True,
    offset: int = 0,
    limit: int = 50,
):
    """报销列表：默认只看自己的；持 expense:manage 且 mine=false 看全部。"""
    can_all = (not mine) and ("*" in perms or "expense:manage" in perms)
    data = await ExpenseService(session).list(
        user_id=None if can_all else user.id,
        status=status, month=month, q=q, offset=offset, limit=limit,
    )
    return ExpensePage(**data)


@expense_router.get("/stats", response_model=ExpenseStatsOut)
async def expense_stats(session: Db, _: ExpenseManageDep, month: str | None = None):
    return await ExpenseService(session).stats(month)


# ---- 预算/额度（必须注册在 GET /{eid} 之前，否则 "budgets" 会被当成 eid） ----
@expense_router.get("/budgets", response_model=BudgetPage)
async def list_budgets(
    session: Db, _: ExpenseManageDep,
    status: str | None = None, department_id: int | None = None,
    offset: int = 0, limit: int = 50,
):
    return await BudgetService(session).list(status=status, department_id=department_id, offset=offset, limit=limit)


@expense_router.post("/budgets", response_model=BudgetOut, status_code=201)
async def create_budget(body: BudgetCreate, session: Db, _: ExpenseManageDep):
    return await BudgetService(session).create(body.model_dump())


@expense_router.patch("/budgets/{budget_id}", response_model=BudgetOut)
async def update_budget(budget_id: int, body: BudgetUpdate, session: Db, _: ExpenseManageDep):
    return await BudgetService(session).update(budget_id, body.model_dump(exclude_unset=True))


@expense_router.get("/budgets/quota-check", response_model=BudgetQuotaCheck)
async def check_budget_quota(
    session: Db, _: ExpenseViewDep,
    amount: float = Query(..., gt=0),
    category: str | None = None,
    department_id: int | None = None,
    occur_date: date = Query(...),
):
    """额度预检（expense:view）：填单时给申请人实时提示剩余额度/超额。"""
    return await BudgetService(session).check_quota(amount, category, department_id, occur_date)


@expense_router.get("/export")
async def export_expenses(
    session: Db, _: ExpenseManageDep, month: str | None = None
):
    """导出报销明细为 Excel（可按月筛选）。"""
    data = await ExpenseService(session).list(
        user_id=None, status=None, month=month, q=None, offset=0, limit=5000
    )
    rows = [["单号", "申请人", "类型", "金额", "发生日期", "状态", "说明", "创建时间", "审批意见"]]
    for it in data["items"]:
        rows.append([
            it["code"], it["user_name"], it["claim_type"], it["amount"],
            it["occur_date"].isoformat(), it["status"], it["description"],
            str(it["created_at"]), it["review_comment"],
        ])
    return _xlsx_response(f"expenses-{month or 'all'}.xlsx", {"报销明细": rows})


@expense_router.get("/import/template")
async def download_import_template(session: Db, _: ExpenseViewDep):
    """下载报销导入模板：预设表头与示例行。"""
    rows = [
        ["金额", "日期", "类型", "说明"],
        [100.50, "2026-09-15", "travel", "客户拜访差旅"],
        [50.00, "2026-09-16", "daily", "办公用品采购"],
    ]
    return _xlsx_response("expense_import_template.xlsx", {"导入模板": rows})


@expense_router.post("/import")
async def import_expenses(
    session: Db, user: CurrentUser, _: ExpenseViewDep, file: UploadFile = File(...)
):
    """Excel 批量导入为草稿：表头 列名 金额 / 日期 / 类型 / 说明（顺序不限）。"""
    raw = await file.read()
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValidationError("导入 Excel 需要 openpyxl，请执行 pip install openpyxl") from exc
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValidationError(f"无法解析 Excel 文件：{exc}") from exc
    ws = wb.active
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    if not grid:
        raise ValidationError("表格为空")
    header = [str(c).strip() if c is not None else "" for c in grid[0]]
    idx = {name: i for i, name in enumerate(header)}

    def cell(row, name):
        i = idx.get(name)
        return row[i] if i is not None and i < len(row) else None

    svc = ExpenseService(session)
    created, errors = 0, []
    for n, row in enumerate(grid[1:], start=2):
        amount = cell(row, "金额")
        day = cell(row, "日期")
        if amount in (None, "") and day in (None, ""):
            continue  # 跳过空行
        try:
            if isinstance(day, datetime):
                day = day.date()
            elif isinstance(day, str):
                day = datetime.fromisoformat(day.strip()).date()
            claim = await svc.create(user.id, {
                "amount": float(amount),
                "occur_date": day,
                "claim_type": str(cell(row, "类型") or "daily").strip() or "daily",
                "description": str(cell(row, "说明") or ""),
                "attachments": [],
            })
            created += 1
        except Exception as exc:  # noqa: BLE001
            errors.append(f"第 {n} 行：{exc}")
    return {"created": created, "errors": errors[:20]}


@expense_router.get("/{eid}", response_model=ExpenseOut)
async def get_expense(
    eid: int,
    session: Db,
    user: CurrentUser,
    _: ExpenseViewDep,
    perms: Annotated[list[str], Depends(get_current_permissions)],
):
    claim = await _expense_owned_check(session, eid, user, perms)
    return ExpenseOut(**await ExpenseService(session).decorate(claim))


@expense_router.post("", response_model=ExpenseOut, status_code=201)
async def create_expense(body: ExpenseCreate, session: Db, user: CurrentUser, _: ExpenseViewDep):
    claim = await ExpenseService(session).create(user.id, body.model_dump())
    return ExpenseOut(**await ExpenseService(session).decorate(claim))


@expense_router.patch("/{eid}", response_model=ExpenseOut)
async def update_expense(
    eid: int, body: ExpenseUpdate, session: Db, user: CurrentUser, _: ExpenseViewDep
):
    claim = await ExpenseService(session).update(eid, user, body.model_dump(exclude_unset=True))
    return ExpenseOut(**await ExpenseService(session).decorate(claim))


@expense_router.delete("/{eid}", response_model=Message)
async def delete_expense(eid: int, session: Db, user: CurrentUser, _: ExpenseViewDep):
    await ExpenseService(session).delete(eid, user)
    return Message(message="报销单已删除")


@expense_router.post("/{eid}/submit", response_model=ExpenseOut)
async def submit_expense(eid: int, session: Db, user: CurrentUser, _: ExpenseViewDep):
    """提交报销：以预置的 expense-claim 流程发起审批实例。"""
    claim = await ExpenseService(session).submit(eid, user)
    return ExpenseOut(**await ExpenseService(session).decorate(claim))


@expense_router.post("/{eid}/paid", response_model=ExpenseOut)
async def pay_expense(
    eid: int, body: ExpensePaidIn, session: Db, user: CurrentUser, _: ExpenseManageDep
):
    """标记打款 / 撤销打款（仅 expense:manage）。"""
    claim = await ExpenseService(session).mark_paid(eid, user, body.model_dump())
    return ExpenseOut(**await ExpenseService(session).decorate(claim))


# ---- 发票验真 ----
@expense_router.post("/invoice/verify", response_model=InvoiceVerifyResponse)
async def verify_invoice(body: InvoiceVerifyRequest, session: Db, _: ExpenseViewDep):
    """发票验真：格式校验 + 唯一性检测 + 买方税号校验位。"""
    return await InvoiceService(session).verify_invoice(
        invoice_no=body.invoice_no,
        invoice_code=body.invoice_code,
        amount=body.amount,
        buyer_tax_no=body.buyer_tax_no,
    )


# ---- 批量打款 ----
@expense_router.post("/batch-pay", response_model=BatchPayResponse)
async def batch_pay_expenses(
    body: BatchPayRequest, session: Db, user: CurrentUser, _: ExpenseManageDep
):
    """批量打款：一次性处理多个已通过的报销单。"""
    return await BatchPaymentService(session, event_bus=get_event_bus()).batch_pay(
        body.expense_ids, user
    )


# =========================================================================
# ---------- 公文（收发文 / 传阅 / 归档） ----------
# =========================================================================
official_router = APIRouter(prefix="/api/official-docs", tags=["official-docs"])
OfficialViewDep = Annotated[None, Depends(require_permissions("doc:official:view"))]
OfficialManageDep = Annotated[None, Depends(require_permissions("doc:official:manage"))]


@official_router.get("", response_model=OfficialDocPage)
async def list_official_docs(
    session: Db,
    _: OfficialViewDep,
    direction: str | None = None,
    status: str | None = None,
    q: str | None = None,
    offset: int = 0,
    limit: int = 50,
):
    data = await OfficialDocService(session).list(
        direction=direction, status=status, q=q, offset=offset, limit=limit
    )
    return OfficialDocPage(**data)


# ---- 静态路径必须注册在 GET /{did} 之前，否则 "search"/"next-doc-no" 会被当成 did ----
@official_router.get("/next-doc-no")
async def preview_next_doc_no(
    session: Db, _: OfficialViewDep,
    direction: str = Query("in", pattern="^(in|out)$"),
):
    """预览下一个文号（不创建）：供前端实时显示编号。"""
    return await OfficialDocService(session).preview_doc_no(direction)


@official_router.get("/search")
async def search_official_docs(
    session: Db, _: OfficialViewDep,
    q: str = Query(..., min_length=1, max_length=100),
    direction: str | None = None,
    status: str | None = None,
    offset: int = 0,
    limit: int = 20,
):
    """全文检索：跨标题/内容/文号/来源单位/目标单位搜索。"""
    return await OfficialDocService(session).search(q, direction, status, offset, limit)


@official_router.get("/{did}", response_model=OfficialDocOut)
async def get_official_doc(did: int, session: Db, _: OfficialViewDep):
    return await OfficialDocService(session).get(did)


@official_router.post("", response_model=OfficialDocOut, status_code=201)
async def create_official_doc(
    body: OfficialDocCreate, session: Db, user: CurrentUser, _: OfficialManageDep
):
    return await OfficialDocService(session).create(user, body.model_dump())


@official_router.patch("/{did}", response_model=OfficialDocOut)
async def update_official_doc(
    did: int, body: OfficialDocUpdate, session: Db, user: CurrentUser, _: OfficialManageDep
):
    return await OfficialDocService(session).update(did, user, body.model_dump(exclude_unset=True))


@official_router.delete("/{did}", response_model=Message)
async def delete_official_doc(did: int, session: Db, user: CurrentUser, _: OfficialManageDep):
    await OfficialDocService(session).delete(did, user)
    return Message(message="公文已删除")


@official_router.post("/{did}/circulate", response_model=OfficialDocOut)
async def circulate_official_doc(
    did: int, body: OfficialDocCirculate, session: Db, user: CurrentUser, _: OfficialManageDep
):
    """传阅：通知指定用户（为空则全体），状态置 circulated。"""
    return await OfficialDocService(session).circulate(user, did, body.model_dump())


@official_router.post("/{did}/archive", response_model=OfficialDocOut)
async def archive_official_doc(
    did: int, session: Db, user: CurrentUser, _: OfficialManageDep
):
    return await OfficialDocService(session, event_bus=get_event_bus()).archive(user, did)


@official_router.get("/{did}/timeline")
async def official_doc_timeline(did: int, session: Db, _: OfficialViewDep):
    """获取公文流转时间线：创建→传阅→归档等操作记录。"""
    return await OfficialDocService(session).get_timeline(did)


# =========================================================================
# ---------- 文档版本（挂在既有 doc_router / ann_router 上） ----------
# =========================================================================
@doc_router.get("/{doc_id}/versions", response_model=list[DocumentVersionOut])
async def list_document_versions(doc_id: int, session: Db, _: DocViewDep):
    items = await DocumentService(session).list_versions(doc_id)
    out = []
    for v in items:
        out.append(DocumentVersionOut(
            id=v.id, document_id=v.document_id, version=v.version,
            file_type=v.file_type, file_size=v.file_size,
            file_url=f"/api/documents/{doc_id}/versions/{v.id}/file",
            note=v.note or "", uploader_name=v.uploader_name or "",
            created_at=v.created_at,
        ))
    return out


@doc_router.post("/{doc_id}/versions", response_model=DocumentVersionOut, status_code=201)
async def add_document_version(
    doc_id: int, session: Db, user: CurrentUser, _: DocUploadDep,
    note: str = Form(""), file: UploadFile = File(...),
):
    """上传新版本：旧版本保留可回溯下载，文档当前版本随之切换。"""
    v = await DocumentService(session).add_version(doc_id, user, file, note)
    return DocumentVersionOut(
        id=v.id, document_id=v.document_id, version=v.version,
        file_type=v.file_type, file_size=v.file_size,
        file_url=f"/api/documents/{doc_id}/versions/{v.id}/file",
        note=v.note or "", uploader_name=v.uploader_name or "",
        created_at=v.created_at,
    )


@doc_router.get("/{doc_id}/versions/{version_id}/file")
async def download_document_version(doc_id: int, version_id: int, session: Db, _: DocViewDep):
    """下载指定历史版本（本地流式 / S3 签名 URL 302）。"""
    local_path, signed_url, download_name = await DocumentService(session).get_version_file(
        doc_id, version_id
    )
    if signed_url:
        return RedirectResponse(signed_url, status_code=302)
    return FileResponse(local_path, filename=download_name)


@ann_router.post("/{ann_id}/read", response_model=AnnouncementReadOut)
async def mark_announcement_read(ann_id: int, session: Db, user: CurrentUser, _: AnnViewDep):
    """标记公告已读（幂等），回写阅读人数。"""
    return AnnouncementReadOut(**await AnnouncementService(session).mark_read(ann_id, user.id))


@ann_router.get("/{ann_id}/readers", response_model=AnnouncementReadersOut)
async def announcement_readers(ann_id: int, session: Db, _: AnnManageDep):
    """公告阅读明细（发布者/管理员可见）。"""
    return AnnouncementReadersOut(**await AnnouncementService(session).readers(ann_id))


# 汇总
from app.api.realtime import router as realtime_router

api.include_router(auth_router)
api.include_router(msg_router)
api.include_router(user_router)
api.include_router(dept_router)
api.include_router(leave_router)
api.include_router(ann_router)
api.include_router(wf_router)
api.include_router(sig_router)
api.include_router(perm_router)
api.include_router(menu_router)
api.include_router(portal_router)
api.include_router(dashboard_router)
api.include_router(ops_router)
api.include_router(audit_router)
api.include_router(model_router)
api.include_router(doc_router)
api.include_router(attendance_router)
api.include_router(expense_router)
api.include_router(official_router)
api.include_router(realtime_router)
